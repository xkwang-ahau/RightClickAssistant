#!/usr/bin/env python3
"""Generate the app's blank Office resources; never used at app runtime."""

from datetime import datetime, timezone
from io import BytesIO
from importlib.metadata import distribution
from pathlib import Path
import hashlib
import json
import re
import zipfile

import docx
import openpyxl
import pptx
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "RightClickAssistant" / "Resources" / "Templates"
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def clear_core_properties(properties):
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = ""
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME


def canonical_package(buffer):
    """Normalize package timestamps so the same inputs produce stable resources."""
    output = BytesIO()
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as package:
        for name in sorted(original.namelist()):
            data = original.read(name)
            if name == "docProps/core.xml":
                # openpyxl updates modified at save time; normalize both date fields.
                for field in (b"created", b"modified"):
                    data = re.sub(
                        rb"(<dcterms:" + field + rb"\b[^>]*>).*?(</dcterms:" + field + rb">)",
                        rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                        data,
                    )
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            package.writestr(info, data)
    return output.getvalue()


def save_template(document, name):
    raw = BytesIO()
    document.save(raw)
    raw.seek(0)
    data = canonical_package(raw)
    (OUTPUT / name).write_bytes(data)
    return {"file": name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)

    word = Document()
    word.add_paragraph()
    clear_core_properties(word.core_properties)

    slides = Presentation()
    slides.slide_width = Inches(13.333333)
    slides.slide_height = Inches(7.5)
    slides.slides.add_slide(slides.slide_layouts[6])
    clear_core_properties(slides.core_properties)

    workbook = Workbook()
    workbook.active.title = "工作表1"
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME

    records = [
        save_template(word, "BlankWord.docx"),
        save_template(slides, "BlankPowerPoint.pptx"),
        save_template(workbook, "BlankExcel.xlsx"),
    ]
    manifest = {
        "generator": "scripts/generate_office_templates.py",
        "versions": {
            "python-docx": docx.__version__,
            "python-pptx": pptx.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "templates": records,
    }
    (ROOT / "docs" / "office-template-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    notices = ["Office template generation tools and license notices\n"]
    for name in manifest["versions"]:
        package = distribution(name)
        candidates = [
            file for file in package.files
            if file.name.lower() in ("license", "license.txt", "licence.rst")
        ]
        if not candidates:
            raise RuntimeError(f"No license notice found for {name}")
        notices.append(f"\n{name} {package.version}\n{package.locate_file(candidates[0]).read_text()}\n")
    (OUTPUT.parent / "OfficeTemplateLicenses.txt").write_text("\n".join(notices))
    for record in records:
        print(f"Generated {record['file']}: {record['bytes']} bytes")


if __name__ == "__main__":
    main()
