#!/usr/bin/env python3
"""Verify package integrity, relationships, blank content and recorded digests."""

from pathlib import Path
from urllib.parse import unquote
import argparse
import hashlib
import json
import posixpath
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def verify_package(path, record):
    data = path.read_bytes()
    assert hashlib.sha256(data).hexdigest() == record["sha256"], f"Digest mismatch: {path}"
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None, f"CRC failure: {path}"
        names = set(package.namelist())
        assert "[Content_Types].xml" in names and "_rels/.rels" in names
        for name in names:
            if name.endswith((".xml", ".rels")):
                tree = ET.fromstring(package.read(name))
                if name.endswith(".rels"):
                    base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
                    for relationship in tree.findall(f"{REL_NS}Relationship"):
                        assert relationship.get("TargetMode") != "External", f"External link: {name}"
                        target = unquote(relationship.attrib["Target"])
                        resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                        assert resolved in names, f"Broken relationship: {name} -> {resolved}"

    if path.suffix == ".docx":
        document = Document(path)
        assert len(document.paragraphs) == 1
        assert not any(paragraph.text for paragraph in document.paragraphs)
        assert not document.tables
    elif path.suffix == ".pptx":
        slides = Presentation(path)
        assert len(slides.slides) == 1 and len(slides.slides[0].shapes) == 0
    elif path.suffix == ".xlsx":
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["工作表1"]
        assert workbook.active.max_row == 1 and workbook.active.max_column == 1
        assert workbook.active["A1"].value is None
        workbook.close()
    print(f"Verified {path.name}: digest, ZIP/XML, relationships, library loading, blank content")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", type=Path, default=ROOT / "RightClickAssistant" / "Resources" / "Templates")
    args = parser.parse_args()
    manifest = json.loads((ROOT / "docs" / "office-template-manifest.json").read_text())
    for record in manifest["templates"]:
        verify_package(args.directory / record["file"], record)


if __name__ == "__main__":
    main()
